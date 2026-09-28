#!/usr/bin/env python3
# IndexSelect case: migrate the algorithm-visualization tables from a reference
# deck (one slide per SIMT variant) into per-variant single-slide decks, then
# restructure each slide: drop the original title and code block, add a Chinese
# title, put the measured-verdict panel right under the title, shift the
# original diagram below it, and append a takeaway note.
#
# Usage: python cases/index-select-simt/migrate_ref_figures.py <参考.pptx>
# Outputs out/<name>.pptx next to this script. The reference deck is private
# material and is NOT shipped with the repo.
import shutil
import sys
from pathlib import Path

HERE = Path(__file__).resolve().parent
sys.path.insert(0, str(HERE.parents[1] / "scripts"))
from pptx import Presentation
from pptx.oxml.ns import qn
from pptx.util import Emu

from seeten_draw import note, panel, text

OUT = HERE / "out"

TITLES = {
    "gather_addressing": "IndexSelect：离散行 gather 的寻址结构",
    "variant_basic": "basic：每个线程负责一整行",
    "variant_coalesced": "coalesced：每个 block 负责一行（4B 步进）",
    "variant_16B": "16B：每个 block 负责一行（float4 步进）",
    "variant_seg16B": "coalesced_seg_16B：每个 warp 负责行内一段",
    "variant_stream": "16B_stream：输出展平后的全局顺序流",
}

CODE_MARKERS = ("//", "int ", "for (", "__ldg", "const int", "if (rid")

VERDICTS = {
    "variant_basic": (
        [["合并效果", "无：32 条 lane 各奔一行，warp 内无合并，每 lane 4B 逐元素发射"],
         ["实测带宽", "16~212 GB/s，任何形状都是全场最差"]],
        "要点：并行度有了，但 warp 的访存完全离散——行 gather 里最差的数据排布。"),
    "variant_coalesced": (
        [["合并效果", "有：warp 每步 32×4B = 128B 连续事务；blockDim 越大行内推进越快"],
         ["实测带宽", "行内家族基准；K≥4096 后被 16B 版全面取代"]],
        "要点：把 warp 收进同一行是第一个质变；下一步质变是把每笔从 4B 加宽到 16B。"),
    "variant_16B": (
        [["合并效果", "有：warp 每步 32×16B = 512B 连续，事务笔数降为 4B 版的 1/4"],
         ["实测带宽", "行宽 ≥1024 时的行内家族主力"]],
        "要点：行内搬运的 SIMT 上限形态；与连续拷贝 16B 版相同的每笔宽度。"),
    "variant_seg16B": (
        [["合并效果", "有：同为 512B/步；段内 512B 对齐、warp 间无交错"],
         ["实测带宽", "与 16B 全程 ±5% 内；仅 K=8192 时以 <5% 成为行内家族最优"]],
        "要点：分段不带来系统性收益，只在行很宽时略优——列入负收益优化清单。"),
    "variant_stream": (
        [["合并效果", "完全顺序（同连续拷贝）；每 16B 付一次除法 + 一次 index 读取"],
         ["实测带宽", "K≤1024 的全场最优 SIMT；K≥4096 被行内家族反超"]],
        "要点：小行宽时把 warp 排进连续地址的最简方案；行变宽后除法与 index 重读成为净开销。"),
}

GATHER_FACTS = (
    [["维度", "表现"],
     ["行内", "K 个 float 地址连续，可 16B 向量化（同连续拷贝）"],
     ["行间", "行首地址由 index[i] 决定，行与行之间不连续、无合并"],
     ["index", "每个输出行读取一次行号，共 N 个 int"]],
    "要点：gather 的成本 = 行内连续拷贝 + 每行一次离散行首定位；"
    "行间离散部分无法用向量化消除。")


def _inch(v):
    return Emu(v).inches


def isolate(src, keep_idx, out_path):
    shutil.copy(src, out_path)
    prs = Presentation(str(out_path))
    sldIdLst = prs.slides._sldIdLst
    for i, sldId in enumerate(list(sldIdLst)):
        if i != keep_idx:
            prs.part.drop_rel(sldId.get(qn("r:id")))
            sldIdLst.remove(sldId)
    return prs, prs.slides[0]


def strip_title_and_code(s):
    for sh in list(s.shapes):
        txt = sh.text_frame.text if sh.has_text_frame else ""
        if sh.top is not None and _inch(sh.top) < 2.2 and any(m in txt for m in CODE_MARKERS):
            sh._element.getparent().remove(sh._element)
    for sh in list(s.shapes):
        txt = (sh.text_frame.text or "").strip() if sh.has_text_frame else ""
        if txt and sh.top is not None and _inch(sh.top) < 1.3 and _inch(sh.width) > 3:
            sh._element.getparent().remove(sh._element)
            break


def shift_all(s, delta_in):
    d = int(delta_in * 914400)
    for sh in s.shapes:
        if sh.top is not None:
            sh.top = Emu(int(sh.top) + d)


def bounds(s):
    tops, bots = [], []
    for sh in s.shapes:
        if sh.top is not None and sh.height is not None:
            tops.append(_inch(sh.top))
            bots.append(_inch(sh.top) + _inch(sh.height))
    return min(tops), max(bots)


def build(src, name, keep_idx, header, rows, take, target_top, col_w):
    out = OUT / f"{name}.pptx"
    prs, s = isolate(src, keep_idx, out)
    strip_title_and_code(s)
    shift_all(s, target_top - bounds(s)[0])   # shift the ORIGINAL diagram first
    text(s, 0.45, 0.28, TITLES[name], w=12.4, size=22.0, bold=True)
    bot = panel(s, 0.45, 1.0,
                "实测定位" if name.startswith("variant") else "寻址事实",
                header, rows, col_w=col_w)
    note(s, 0.45, bot + 0.12, take, w=12.4)
    prs.slide_height = Emu(int((bounds(s)[1] + 0.3) * 914400))
    prs.save(str(out))
    print(f"{name} -> {out}")


def main():
    if len(sys.argv) != 2:
        sys.exit("usage: migrate_ref_figures.py <参考.pptx>")
    src = sys.argv[1]
    OUT.mkdir(exist_ok=True)
    header, rows = GATHER_FACTS[0][0], GATHER_FACTS[0][1:]
    build(src, "gather_addressing", 0, header, rows, GATHER_FACTS[1], 3.75,
          col_w=[1.3, 11.0])
    for name, keep_idx in [("variant_basic", 2), ("variant_coalesced", 3),
                           ("variant_16B", 4), ("variant_seg16B", 5),
                           ("variant_stream", 6)]:
        rows, take = VERDICTS[name]
        build(src, name, keep_idx, ["项", "结论"], rows, take, 3.35,
              col_w=[1.7, 10.6])


if __name__ == "__main__":
    main()
