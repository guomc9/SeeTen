"""对生成出来的 pptx 做内容自检 —— 公式回代 + 矩阵回读 + 结论核对。

    python check_case.py out/v3-index-schedules.pptx

四类检查（对应 style-spec.md §7 的内容正确性约定）：

1. 不变量：七个例子的 `check()` 全过（任务不重复 / 同轮 S1 不撞 / 列私有）；
2. 矩阵回读：从 pptx 里读回「轮 \\ 核」任务矩阵，与算法函数逐格对照；
3. 公式回代：把**页面上印着的伪代码公式**重算一遍，与算法函数输出逐格对照 ——
   公式、数据、页面三者必须一致；
4. 结论核对：把几处容易写歪的结论（GQA 是否保证列私有、Left-Up m>n 是否 idle、
   对比表"用不上"行有没有给出本例不满足的条件）用数据判定后与页面文案对照。

退出码 0 = 全部通过。
"""

from __future__ import annotations

import os
import sys

HERE = os.path.dirname(os.path.abspath(__file__))
sys.path.insert(0, os.path.join(HERE, "..", "..", "scripts"))
sys.path.insert(0, HERE)

from pptx import Presentation                    # noqa: E402

import draw_case as dc                            # noqa: E402
import index_schedules as ix                      # noqa: E402
import seeten_draw as sd                          # noqa: E402

FAILURES = []


def fail(msg):
    FAILURES.append(msg)
    print(f"  [FAIL] {msg}")


def ok(msg):
    print(f"  [ OK ] {msg}")


# ------------------------------------------------------- 公式回代（页面公式）

def f_dense_swizzle(r, j, k, m, n, b):
    """slide 3 印刷公式（1-based）。"""
    p = (r - 1) // m * k + (j - 1)
    col = p % n + 1
    bat = p // n + 1
    row = (col - 1 + r - 1) % m + 1
    if not (1 <= bat <= b and 1 <= col <= n and 1 <= row <= m):
        return None
    return (bat - 1, 0, 0, row - 1, col - 1)


def f_dense_index(r, j, k, m, n, b):
    """slide 5 印刷公式（1-based）。"""
    p = (r - 1) // m * k + j
    bat = p % b or b
    col = -(-p // b)
    row = ((col - 1) % m + (r - 1) % m) % m + 1
    if not (1 <= bat <= b and 1 <= col <= n and 1 <= row <= m):
        return None
    return (bat - 1, 0, 0, row - 1, col - 1)


def f_causal_swizzle(r, j, k, m, n, b):
    """slide 8 印刷公式：虚拟 m x (n+1) 里先跑 Dense Swizzle，再折回真实坐标。"""
    v = f_dense_swizzle(r, j, k, m, n + 1, b >> 1)
    if v is None:
        return None
    bat, x, y = v[0] + 1, v[3] + 1, v[4] + 1
    if y >= x + 1:
        y, x, bat = (n << 1) - m - y + 2, m + 1 - x, bat << 1
    else:
        bat = (bat << 1) - 1
    if not (1 <= bat <= b and 1 <= x <= m and 1 <= y <= n):
        return None
    return (bat - 1, 0, 0, x - 1, y - 1)


def f_gqa_dense(r, j, k, m, n, b, g):
    """slide 13 印刷公式（本例不触发 t1*m < n 的 ID 重排修正）。"""
    R = max(-(-b * n * g // k), -(-n // m), g)
    ID = (j - 1) * R + -(-r // m)
    bat = -(-ix.nz(ID % (b * g), b * g) // g)
    col = -(-ID // (b * g))
    grp = ix.nz(ID % g, g)
    gd = ix.gcd(b * g, R)
    t1 = R // gd
    off = -(-ix.nz(col % (t1 * m), t1 * m) // t1)
    row = ix.nz(r % m, m) + off - 1
    if row > m:
        row -= m
    return (bat - 1, 0, grp - 1, row - 1, col - 1)


def f_tnd_dense(shape, kw, r, j):
    """slide 15 印刷公式（批内 0-based）。"""
    for b in range(shape.batch):
        if not (kw["prefix"][b] < r <= kw["prefix"][b + 1]):
            continue
        delta = r - kw["prefix"][b] - 1
        s1o, s2o = dc.batch_mn(shape, kw, b)
        tid = delta // s1o * shape.coreNum + (j - 1)
        if tid >= s2o:
            return None
        col = tid % s2o
        row = (col + delta) % s1o
        return (b, 0, 0, row, col)
    return None


def f_tnd_gqa(shape, kw, mr, r, j):
    """slide 17 印刷公式（本例不触发 t1 < n 的 gcd 对齐修正）。"""
    n1 = shape.kvHeadNum * shape.groupNum
    ID = (j - 1) * mr + r
    w = 0
    while w + 1 < shape.batch and ID > kw["area_prefix"][w + 1] * n1:
        w += 1
    if ID > kw["area_prefix"][w + 1] * n1:
        return None
    m, n = dc.batch_mn(shape, kw, w)
    g = shape.groupNum
    if mr // ix.gcd(m * g, mr) < n:
        return None                      # 触发修正的 shape 不走本页简化公式
    inner = ID - kw["area_prefix"][w] * n1
    piece = -(-inner // (m * n * g))
    within = ix.nz(inner % (m * n * g), m * n * g)
    col = -(-within // (m * g))
    grp = -(-ix.nz(within % (m * g), m * g) // m)
    row = ix.nz(ix.nz(within % (m * g), m * g) % m, m)
    return (w, piece - 1, grp - 1, row - 1, col - 1)


# ------------------------------------------------------- 辅助

def slide_of(kind):
    """算法页（『轮 \\ 核』矩阵所在页）的 slide 序号。"""
    idx = 3                                   # 1 总览 + 2 轴遍历
    for _, k, _, _, _, _ in ix.cases():
        if k == ix.KIND_CAUSAL_SWIZZLE:
            idx += 1                          # 折叠前插一页虚拟坐标系
        if k == kind:
            return idx
        idx += 2                              # 算法页 + 例子页
        if k == ix.KIND_LEFT_UP_CAUSAL:
            idx += 1                          # 5b 非方形页
    raise KeyError(kind)


def api(shape, g):
    return dict(shape=shape, g=g)


def replay(tag, fn, kind, shape, mr, kw, g=None):
    """把 fn(r, j) 公式回代，与算法函数 decode() 逐格对照。"""
    bad = 0
    for r in range(1, mr + 1):
        for j in range(1, shape.coreNum + 1):
            got = fn(r, j)
            ref = ix.decode(kind, shape, r, j, **kw)
            ref_t = None if ref is None else (ref.batch, ref.n2, ref.g,
                                              ref.s1, ref.s2)
            if got != ref_t:
                bad += 1
                print(f"       {tag} r={r} j={j}: 公式={got} 数据={ref_t}")
    if bad:
        fail(f"{tag}: 公式回代 {bad} 格失配")
    else:
        ok(f"{tag}: 公式回代 0 失配")


def matrix_text(kind, shape, c):
    lines = [f"S1={c.s1 + 1}  S2={c.s2 + 1}"]
    if shape.groupNum > 1 or shape.kvHeadNum > 1:
        lines.append(f"B={c.batch + 1} N2={c.n2 + 1} G={c.g + 1}")
    elif shape.batch > 1:
        lines.append(f"B={c.batch + 1}")
    return " ".join(lines)


def expected_matrix(kind, shape, mr, kw):
    ts = ix.tasks_of(kind, shape, mr, **kw)
    rows = [[f"轮 \\ 核"] + [f"C{j}" for j in range(1, shape.coreNum + 1)]]
    for r in range(1, mr + 1):
        row = [f"第 {r} 轮"]
        for j in range(1, shape.coreNum + 1):
            c = ts.get((j, r))
            row.append("" if c is None else matrix_text(kind, shape, c))
        rows.append(row)
    return rows


# ------------------------------------------------------- 主流程

def main(pptx_path):
    print("== 1. 算法不变量（七个例子 + 5b）")
    for name, kind, shape, mr, kw, _ in ix.cases():
        r = ix.check(kind, shape, mr, **kw)
        bad = r["dup"] + r["dq_conflict"]
        need_priv = kind not in (ix.KIND_GQA_DENSE, ix.KIND_TND_GQA_DENSE)
        priv_txt = r["column_private"] if need_priv else "不适用(GQA)"
        if bad or (need_priv and not r["column_private"]):
            fail(f"{name}: dup={r['dup']} dq={r['dq_conflict']} "
                 f"列私有={priv_txt}")
        else:
            ok(f"{name}: dup=0 dq=0 列私有={priv_txt} idle={r['holes']}")

    s5b = ix.Shape(batch=2, qSeqLen=4 * 128, kvSeqLen=3 * 128, qHeadNum=1,
                   kvHeadNum=1, coreNum=3)
    r5b = ix.check(ix.KIND_LEFT_UP_CAUSAL, s5b, 6)
    if r5b["holes"] or not r5b["column_private"]:
        fail(f"5b 形状: holes={r5b['holes']} 列私有={r5b['column_private']}")
    else:
        ok("5b 形状: 18/18 槽位零 idle、列私有")

    print("== 2. pptx 矩阵回读 vs 算法函数")
    prs = Presentation(pptx_path)
    expected = {}
    for _, kind, shape, mr, kw, _ in ix.cases():
        expected[slide_of(kind)] = expected_matrix(kind, shape, mr, kw)
    mism = sd.check_task_matrices(prs, expected)
    if mism:
        for sl, r, c, e, got in mism[:10]:
            print(f"       slide {sl} r{r}c{c}: 期望「{e}」实际「{got}」")
        fail(f"矩阵回读: {len(mism)} 格失配")
    else:
        ok(f"矩阵回读: {sum(len(v) for v in expected.values())} 行全部一致")

    print("== 3. 印刷公式回代")
    s1 = ix.Shape(2, 384, 384, 1, 1, coreNum=2)
    s2 = ix.Shape(2, 256, 256, 1, 1, coreNum=4)
    s3 = ix.Shape(2, 384, 384, 1, 1, coreNum=2)
    s5 = ix.Shape(1, 256, 256, 2, 1, groupNum=2, coreNum=2)
    s6 = ix.Shape(batch=2, qHeadNum=1, kvHeadNum=1, coreNum=2)
    kw6 = dict(cu_q=[384, 640], cu_k=[256, 640], prefix=[0, 3, 7])
    s7 = ix.Shape(batch=2, qHeadNum=2, kvHeadNum=1, groupNum=2, coreNum=2)
    kw7 = dict(cu_q=[256, 384], cu_k=[256, 512], area_prefix=[0, 4, 6])
    replay("Dense Swizzle (p3)", lambda r, j: f_dense_swizzle(r, j, 2, 3, 3, 2),
           ix.KIND_DENSE_SWIZZLE, s1, 9, {})
    replay("Dense Index (p5)", lambda r, j: f_dense_index(r, j, 4, 2, 2, 2),
           ix.KIND_DENSE_INDEX, s2, 2, {})
    replay("Causal Swizzle (p8)", lambda r, j: f_causal_swizzle(r, j, 2, 3, 3, 2),
           ix.KIND_CAUSAL_SWIZZLE, s3, 6, {})

    # Left-Up 方形页 = 委托方形折叠：与 Causal Swizzle 逐格相同
    same = all(f_causal_swizzle(r, j, 2, 3, 3, 2) ==
               (lambda c: None if c is None else (c.batch, c.n2, c.g, c.s1, c.s2))
               (ix.decode(ix.KIND_LEFT_UP_CAUSAL, s3, r, j))
               for r in range(1, 7) for j in range(1, 3))
    (ok if same else fail)("Left-Up 方形页 = 委托方形折叠（逐格相同）")

    replay("GQA Dense (p13)", lambda r, j: f_gqa_dense(r, j, 2, 2, 2, 1, 2),
           ix.KIND_GQA_DENSE, s5, 4, {})
    replay("TND Dense (p15)", lambda r, j: f_tnd_dense(s6, kw6, r, j),
           ix.KIND_TND_DENSE, s6, 7, kw6)
    replay("TND GQA (p17)", lambda r, j: f_tnd_gqa(s7, kw7, 6, r, j),
           ix.KIND_TND_GQA_DENSE, s7, 6, dict(kw7, max_round=6))

    print("== 4. 结论句核对")
    # GQA：本案例的列是否跨核 —— 文案必须按"不保证"表述
    owners = {}
    for r in range(1, 5):
        for j in range(1, 3):
            c = ix.decode(ix.KIND_GQA_DENSE, s5, r, j)
            if c is not None:
                owners.setdefault((c.batch, c.n2, c.s2), set()).add(j)
    max_cores = max(len(v) for v in owners.values())
    txt = dc.page_text(ix.KIND_GQA_DENSE, s5, 4, {})
    if max_cores == 1 and "不保证" in txt["key"]:
        ok("GQA: 例子每列同核，文案按『不保证同核』表述")
    else:
        fail(f"GQA: 例子每列 {max_cores} 核，文案需与之一致（key=「{txt['key']}」）")

    # 对比表：用不上的行必须给出"本例不满足"的原因
    for kind, shape, mr, kw in ((ix.KIND_CAUSAL_SWIZZLE, s3, 6, {}),
                                (ix.KIND_GQA_DENSE, s5, 4, {}),
                                (ix.KIND_TND_DENSE, s6, 7, kw6)):
        _, rows = dc.applicability_rows(kind, shape, mr, kw)
        bad = [(cn, st, cd) for cn, st, cd in rows
               if st == "用不上" and not cd.startswith("本例")]
        if bad:
            fail(f"对比表({kind}): 用不上但没说明原因 {bad[:2]}")
        else:
            ok(f"对比表({kind}): 用不上行均给出本例原因")

    # 可达性：Causal Swizzle / Left-Up 必须带可达性标签，5b 页必须标注不可达
    for kind in (ix.KIND_CAUSAL_SWIZZLE, ix.KIND_LEFT_UP_CAUSAL):
        tags = [k for k, v in dc.META[kind]["tags"](s3, {})]
        if "可达性" not in tags:
            fail(f"META[{kind}]: 缺可达性标签")
        else:
            ok(f"META[{kind}]: 可达性 = "
               f"{dict(dc.META[kind]['tags'](s3, {}))['可达性']}")

    src = open(os.path.join(HERE, "draw_case.py"), encoding="utf-8").read()
    if "选择器不会走这一支" not in src:
        fail("5b 页缺『选择器不会走这一支』标注")
    else:
        ok("5b 页已标注选择器不可达")

    # 性能对比页（4 页）：6 列分组统计表 + 明细表逐格 + 口径标注
    import math as _m

    import perf_matrix as pm

    SIZES = ("小", "中", "大")

    def _gm(xs):
        xs = [x for x in xs if x is not None and x > 0]
        return _m.exp(sum(_m.log(x) for x in xs) / len(xs)) if xs else float("nan")

    def _ratio(num, den):
        return [a / b if (a and b) else None for a, b in zip(num, den)]

    def _gvals(vals, lay, sz):
        return [vals[i] for i, (n, l, _) in enumerate(pm.CASES)
                if l == lay and dc._size_class(n) == sz]

    def _rate(vals, th=0.8):
        xs = [v for v in vals if v is not None]
        return sum(1 for v in xs if v >= th) / len(xs) if xs else 0.0

    def _extreme(vals, mode):
        xs = [v for v in vals if v is not None]
        if not xs:
            return float("nan")
        return min(xs) if mode == "min" else max(xs)

    checks_tables, perf_texts = [], []
    for sl in prs.slides:
        for sh in sl.shapes:
            if sh.has_text_frame:
                perf_texts.append(sh.text_frame.text)
            if getattr(sh, "has_table", False) and sh.has_table:
                checks_tables.append(sh.table)
                perf_texts.extend(tc.text for row in sh.table.rows
                                  for tc in row.cells)

    def find_table(header, labels):
        for t in checks_tables:
            if t.cell(0, 0).text.strip() != header:
                continue
            if ([t.cell(r, 0).text.strip() for r in range(1, len(t.rows))]
                    == labels):
                return t
        return None

    def check_agg(header, rows, fmts, tag):
        """6 列分组统计表：BSND/TND x 小/中/大，按行口径由数据模块重算。"""
        t = find_table(header, [r[0] for r in rows])
        if t is None:
            fail(f"性能页缺统计表: {header} / {[r[0] for r in rows]}")
            return
        fn_of = {"gm": _gm, "rate": _rate,
                 "min": lambda v: _extreme(v, "min"),
                 "max": lambda v: _extreme(v, "max")}
        bad = 0
        for ri, (label, vals, mode) in enumerate(rows):
            fn = fn_of[mode]
            exp = [fmts[ri](fn(_gvals(vals, lay, sz)))
                   for lay in ("BSND", "TND") for sz in SIZES]
            got = [t.cell(ri + 1, ci + 1).text.strip() for ci in range(6)]
            if got != exp:
                bad += 1
                fail(f"性能统计表 {header}/{label}: {got} != {exp}")
        if not bad:
            ok(f"性能统计表 {header}: 与数据一致（{tag}）")

    pen_ours = _ratio(pm.OURS_DET, pm.OURS_ND)
    pen_opst = _ratio(pm.OPST_DET, pm.OPST_ND)
    det_ratio = _ratio(pm.OPST_DET, pm.OURS_DET)
    nd_ratio = _ratio(pm.OPST_ND, pm.OURS_ND)
    f2 = lambda v: f"{v:.2f}"
    fp = lambda v: f"{v:.0%}"
    check_agg("det/nd 开销 (×)",
              [("ours det/nd 几何平均", pen_ours, "gm"),
               ("opst det/nd 几何平均", pen_opst, "gm"),
               ("ours 最差组", pen_ours, "max")],
              [f2, f2, f2], "9.1")
    for header, ratio, tag in (("det 对比 (×)", det_ratio, "9.2"),
                               ("nd 对比 (×)", nd_ratio, "9.3")):
        check_agg(header,
                  [("opst/ours 几何平均", ratio, "gm"),
                   ("达标 (≥0.8×)", ratio, "rate"),
                   ("最差（组内 min）", ratio, "min")],
                  [f2, fp, f2], tag)

    # 9.4 明细：4 张表（BSND 11+10 / TND 8+7），逐格核验
    detail = [t for t in checks_tables
              if t.cell(0, 0).text.strip() == "核时 (μs)"]
    bsnd = [i for i, c in enumerate(pm.CASES) if c[1] == "BSND"]
    tnd = [i for i, c in enumerate(pm.CASES) if c[1] == "TND"]
    blocks = [bsnd[:11], bsnd[11:], tnd[:8], tnd[8:]]
    if len(detail) != len(blocks):
        fail(f"性能明细表数量 {len(detail)} != {len(blocks)}")
    else:
        bad = 0
        for t, idxs in zip(detail, blocks):
            for ri, ci_ in enumerate(idxs):
                exp = [pm.OURS_DET[ci_], pm.OURS_ND[ci_],
                       pm.OPST_DET[ci_], pm.OPST_ND[ci_]]
                exp_s = ["—" if v is None else f"{v:.1f}" for v in exp]
                got = [t.cell(ri + 1, c).text.strip() for c in range(1, 5)]
                if got != exp_s:
                    bad += 1
                    fail(f"明细表行 {pm.CASES[ci_][0]}: {got} != {exp_s}")
        if not bad:
            ok(f"性能明细表: {sum(len(b) for b in blocks)} 行逐格一致")

    blob = " ".join(perf_texts)
    for token in ("非 event record", "msprof", "Task Duration", "det/nd",
                  "opst-det / 本仓库-det", "opst-nd / 本仓库-nd"):
        if token in blob:
            ok(f"性能页标注含「{token}」")
        else:
            fail(f"性能页缺标注「{token}」")

    print()
    if FAILURES:
        print(f"FAILED: {len(FAILURES)} 项")
        return 1
    print("ALL PASS")
    return 0


if __name__ == "__main__":
    sys.exit(main(sys.argv[1] if len(sys.argv) > 1
                  else "out/v3-index-schedules.pptx"))
