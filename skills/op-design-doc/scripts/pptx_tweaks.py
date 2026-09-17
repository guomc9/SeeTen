#!/usr/bin/env python3
"""pptx 小修（用于文档配图前处理）：去标题页号 / 蓝格白字加粗。

用法：
    python pptx_tweaks.py in.pptx out.pptx --strip-title-numbers --bold-white-on-dark
    python pptx_tweaks.py in.pptx out.pptx --pages 2,3,4 --strip-title-numbers

--pages  只处理这些页（1-based）；缺省全部
--strip-title-numbers   删除标题文本里的 "N. " 前缀（判定：run 字号 >= 28pt）
--bold-white-on-dark    深色底 + 白字的表格单元格加粗（渲染更清晰）
"""
import argparse
import re
import sys

from pptx import Presentation


def is_dark(rgb_hex):
    f = str(rgb_hex).lstrip("#")
    if len(f) != 6:
        return False
    r, g, b = int(f[0:2], 16), int(f[2:4], 16), int(f[4:6], 16)
    return (0.299 * r + 0.587 * g + 0.114 * b) < 150


def run_color(r):
    try:
        if r.font.color and r.font.color.type is not None:
            return str(r.font.color.rgb)
    except Exception:
        pass
    return None


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("src")
    ap.add_argument("dst")
    ap.add_argument("--pages")
    ap.add_argument("--strip-title-numbers", action="store_true")
    ap.add_argument("--bold-white-on-dark", action="store_true")
    a = ap.parse_args()

    prs = Presentation(a.src)
    pages = ([int(x) - 1 for x in a.pages.split(",")] if a.pages
             else list(range(len(prs.slides._sldIdLst))))
    n_title = n_bold = 0
    for si in pages:
        slide = prs.slides[si]
        for sh in slide.shapes:
            if a.strip_title_numbers and sh.has_text_frame:
                for p in sh.text_frame.paragraphs:
                    for r in p.runs:
                        if r.font.size and r.font.size.pt >= 28:
                            # 支持 "2." 与 "2.1." 两种页号前缀
                            new = re.sub(r"^(?:\d+\.)+\s*", "", r.text)
                            if new != r.text:
                                r.text = new
                                n_title += 1
            if a.bold_white_on_dark and getattr(sh, "has_table", False) and sh.has_table:
                for row in sh.table.rows:
                    for cell in row.cells:
                        dark = False
                        try:
                            if cell.fill.type is not None:
                                dark = is_dark(cell.fill.fore_color.rgb)
                        except Exception:
                            pass
                        if not dark:
                            continue
                        for p in cell.text_frame.paragraphs:
                            for r in p.runs:
                                if run_color(r) == "FFFFFF" and not r.font.bold:
                                    r.font.bold = True
                                    n_bold += 1
    prs.save(a.dst)
    print(f"{a.dst}: titles={n_title}, bold-runs={n_bold}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
