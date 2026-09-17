#!/usr/bin/env python3
"""pptx → PNG（真渲染：LibreOffice headless + PyMuPDF）。

用法：
    python render_pptx.py deck.pptx --pages 2,3,4 --dpi 140 --outdir figs/
    python render_pptx.py deck.pptx --pages 5 --crop 0.62,1.72,13.95,4.02 --out figs/x.png
    python render_pptx.py deck.pptx --extract-pictures --outdir figs/   # 抽 pptx 内嵌原图

--pages     要导出的页（1-based，逗号分隔）；缺省全部
--crop      x0,y0,x1,y1（英寸），对所有页同一裁剪；与 --out 搭配（单页）
--outdir    输出目录；文件名 = <stem>-p<页号>.png（--pages 单页且给 --out 时用 --out）
--dpi       输出分辨率（缺省 140）
--cache     PDF 缓存目录（缺省 /tmp/op-design-doc-render）

LibreOffice 入口探测顺序：$OP_DESIGN_DOC_LO > 常见路径。
"""
import argparse
import glob
import os
import shutil
import subprocess
import sys

LO_CANDIDATES = [
    os.path.expanduser("~/tools/squashfs-root/AppRun"),
    "/data/g00977778/tools/squashfs-root/AppRun",
    "/opt/libreoffice/squashfs-root/AppRun",
]


def find_lo():
    p = os.environ.get("OP_DESIGN_DOC_LO")
    if p and os.path.exists(p):
        return p
    for c in LO_CANDIDATES:
        if os.path.exists(c):
            return c
    p = shutil.which("soffice") or shutil.which("libreoffice")
    if p:
        return p
    raise SystemExit("找不到 LibreOffice：设置 OP_DESIGN_DOC_LO 或安装便携 AppImage"
                     "（见 references/pptx-render.md）")


def to_pdf(pptx, cache):
    os.makedirs(cache, exist_ok=True)
    stem = os.path.splitext(os.path.basename(pptx))[0]
    pdf = os.path.join(cache, stem + ".pdf")
    if os.path.exists(pdf) and os.path.getmtime(pdf) >= os.path.getmtime(pptx):
        return pdf
    lo = find_lo()
    subprocess.run([lo, "--headless",
                    "-env:UserInstallation=file:///tmp/lo-profile",
                    "--convert-to", "pdf", "--outdir", cache, pptx],
                   check=True, stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL)
    if not os.path.exists(pdf):
        raise SystemExit(f"转换失败：{pdf} 未生成")
    return pdf


def parse_pages(s, n):
    if not s:
        return list(range(n))
    return [int(x) - 1 for x in s.split(",") if x.strip()]


def extract_pictures(pptx, outdir):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE
    os.makedirs(outdir, exist_ok=True)
    prs = Presentation(pptx)
    n = 0
    for si, slide in enumerate(prs.slides, start=1):
        for pi, sh in enumerate(slide.shapes):
            if sh.shape_type == MSO_SHAPE_TYPE.PICTURE:
                path = os.path.join(outdir, f"pic-slide{si}-{pi}.png")
                open(path, "wb").write(sh.image.blob)
                n += 1
                print("wrote", path)
    print(f"{n} pictures extracted")
    return n


def main():
    ap = argparse.ArgumentParser()
    ap.add_argument("pptx")
    ap.add_argument("--pages")
    ap.add_argument("--crop")
    ap.add_argument("--outdir", default=".")
    ap.add_argument("--out")
    ap.add_argument("--dpi", type=float, default=140)
    ap.add_argument("--extract-pictures", action="store_true")
    ap.add_argument("--cache", default="/tmp/op-design-doc-render")
    a = ap.parse_args()

    if a.extract_pictures:
        extract_pictures(a.pptx, a.outdir)
        return 0

    import pymupdf
    pdf = to_pdf(a.pptx, a.cache)
    doc = pymupdf.open(pdf)
    pages = parse_pages(a.pages, doc.page_count)
    clip = None
    if a.crop:
        x0, y0, x1, y1 = (float(v) for v in a.crop.split(","))
        clip = pymupdf.Rect(x0 * 72, y0 * 72, x1 * 72, y1 * 72)
    mat = pymupdf.Matrix(a.dpi / 72, a.dpi / 72)
    stem = os.path.splitext(os.path.basename(a.pptx))[0]
    os.makedirs(a.outdir, exist_ok=True)
    for pi in pages:
        if a.out and len(pages) == 1:
            path = a.out
        else:
            path = os.path.join(a.outdir, f"{stem}-p{pi + 1}.png")
        pix = doc[pi].get_pixmap(matrix=mat, clip=clip)
        pix.save(path)
        print("wrote", path, f"{pix.width}x{pix.height}")
    return 0


if __name__ == "__main__":
    sys.exit(main())
